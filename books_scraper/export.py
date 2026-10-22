"""Write results to CSV, JSON, or Excel."""

from __future__ import annotations

import csv
import json
from pathlib import Path

from .parse import FIELDS, Book

FORMATS = ("csv", "json", "xlsx")


def save(books: list[Book], base: Path, formats: list[str]) -> list[Path]:
    """Save the books in each format next to ``base`` and return the paths."""
    base.parent.mkdir(parents=True, exist_ok=True)
    rows = [b.as_row() for b in books]
    written = []
    for fmt in formats:
        path = base.with_suffix("." + fmt)
        if fmt == "csv":
            with path.open("w", newline="", encoding="utf-8-sig") as fh:
                writer = csv.DictWriter(fh, fieldnames=FIELDS)
                writer.writeheader()
                writer.writerows(rows)
        elif fmt == "json":
            path.write_text(json.dumps(rows, ensure_ascii=False, indent=2), encoding="utf-8")
        elif fmt == "xlsx":
            _save_xlsx(books, path)
        else:
            msg = f"unknown format {fmt!r}"
            raise ValueError(msg)
        written.append(path)
    return written


def _save_xlsx(books: list[Book], path: Path) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    ws.title = "books"
    ws.append(FIELDS)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for b in books:
        ws.append([b.title, float(b.price), b.currency, b.rating, b.in_stock, b.category, b.url])
    for row in ws.iter_rows(min_row=2, min_col=2, max_col=2):
        row[0].number_format = "0.00"
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for letter, width in zip("ABCDEFG", (50, 9, 9, 7, 9, 20, 60), strict=True):
        ws.column_dimensions[letter].width = width
    wb.save(path)
