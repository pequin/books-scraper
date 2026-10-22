import csv
import json

import pytest
from openpyxl import load_workbook

from books_scraper.cli import main
from books_scraper.crawl import crawl


def test_crawl_follows_pagination_within_a_category(site_fetcher):
    books = list(crawl(site_fetcher, only=["mystery"]))

    assert len(books) == 32
    assert {b.category for b in books} == {"Mystery"}
    assert site_fetcher.stats.requests == 3


def test_crawl_rejects_unknown_category(site_fetcher):
    with pytest.raises(ValueError, match="unknown categories: cooking"):
        list(crawl(site_fetcher, only=["Cooking"]))


def test_crawl_skips_a_failing_category_and_keeps_going(site_fetcher):
    # Only Mystery has saved pages; Travel answers 404 and is skipped.
    books = list(crawl(site_fetcher, only=["Travel", "Mystery"]))

    assert len(books) == 32


def test_crawl_respects_max_pages(site_fetcher):
    assert len(list(crawl(site_fetcher, only=["Mystery"], max_pages=1))) == 20


def test_cli_saves_all_formats(tmp_path, site_fetcher, capsys):
    out = tmp_path / "books"

    code = main(["-c", "Mystery", "-o", str(out), "-f", "csv,json,xlsx"], fetcher=site_fetcher)

    assert code == 0
    with (tmp_path / "books.csv").open(encoding="utf-8-sig", newline="") as fh:
        rows = list(csv.DictReader(fh))
    assert len(rows) == 32
    assert rows[0]["title"] == "Sharp Objects"
    assert rows[0]["price"] == "47.82"
    data = json.loads((tmp_path / "books.json").read_text(encoding="utf-8"))
    assert data[0]["rating"] == 4
    sheet = load_workbook(tmp_path / "books.xlsx").active
    assert sheet.max_row == 33
    assert sheet["B2"].value == 47.82
    assert "32 books, 3 requests, 0 retries" in capsys.readouterr().out
